import os
import re
from decimal import Decimal, InvalidOperation
from datetime import datetime, date
from openpyxl import load_workbook


# ============================================================
# FILE PATHS
# ============================================================

EXCEL_PATH = (
    r"C:\testTKC\OpenAI_Invoice_Processing_AP\INPUT\1100\All_Invoices.xlsx"
)

TEXT_PATH = (
    r"C:\testTKC\OpenAI_Invoice_Processing_AP\INPUT\1100\summary invoice report_test.txt"
)

MASTER_PATH = (
    r"C:\testTKC\OpenAI_Invoice_Processing_AP\MasterSuplier"
    r"\Master Vendor Name (MAT)_TH_ENG.xlsx"
)

OUTPUT_EXCEL_PATH = (
    r"C:\testTKC\OpenAI_Invoice_Processing_AP\INPUT\1100"
    r"\OutputExcel\ocr_output_matched.xlsx"
)


# ============================================================
# NORMALIZE FUNCTIONS
# ============================================================

def clean(value):
    if value is None:
        return ""
    return str(value).strip().lstrip("\ufeff")


def normalize_invoice(value):
    return re.sub(r"\s+", "", clean(value)).casefold()


def normalize_tax_id(value):
    return re.sub(r"\D", "", clean(value))


def normalize_text(value):
    return re.sub(
        r"\s+", " ", clean(value)
    ).strip().casefold()


def normalize_branch(value):
    return re.sub(
        r"\s+", "", clean(value)
    ).casefold()


def normalize_date(value):
    """
    Excel format: dd/mm/yyyy
    TXT format: ddmmyyyy
    """

    if value is None or clean(value) == "":
        return ""

    if isinstance(value, (datetime, date)):
        return value.strftime("%d%m%Y")

    text = clean(value)

    for fmt in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d%m%Y",
    ):
        try:
            return datetime.strptime(
                text, fmt
            ).strftime("%d%m%Y")
        except ValueError:
            pass

    return re.sub(r"\D", "", text)


def normalize_amount(value):
    if value is None or clean(value) == "":
        return None

    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(
                Decimal("0.01")
            )
        except InvalidOperation:
            return None

    text = clean(value).replace(",", "")
    text = re.sub(r"[^\d.\-()]", "", text)

    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]

    try:
        return Decimal(text).quantize(
            Decimal("0.01")
        )
    except (InvalidOperation, ValueError):
        return None


# ============================================================
# SUMMARY TXT FUNCTIONS
# ============================================================

def load_summary_h_records(txt_path):
    """
    Read only H records.
    Ignore L and T records.

    Field positions:
        Field 3  = InvoiceDate
        Field 6  = TaxInvoiceNo
        Field 14 = AmountIncVat
        Field 16 = TotalAmount
        Field 17 = VATAmount
        Field 20 = SupplierName + VendorBranch
    """

    records = {}

    with open(
        txt_path,
        "r",
        encoding="utf-8-sig",
        errors="replace"
    ) as file:

        for line_number, line in enumerate(file, 1):

            line = line.rstrip("\r\n")

            # Only process H records.
            if not line.startswith("H"):
                continue

            fields = [
                clean(item)
                for item in line.split("|")
            ]

            if len(fields) < 20:
                continue

            invoice_no = fields[5]
            invoice_key = normalize_invoice(invoice_no)

            if not invoice_key:
                continue

            supplier_field = fields[19]

            # Extract branch from parentheses in Field 20.
            branch_matches = re.findall(
                r"\(([^()]*)\)",
                supplier_field
            )

            vendor_branch = (
                branch_matches[-1].strip()
                if branch_matches else ""
            )

            # Remove trailing branch from supplier name.
            supplier_name = re.sub(
                r"\s*\([^()]*\)\s*$",
                "",
                supplier_field
            ).strip()

            record = {
                "LineNumber": line_number,
                "InvoiceDate": fields[2],
                "TaxInvoiceNo": fields[5],
                "AmountIncVat": fields[13],
                "TotalAmount": fields[15],
                "VATAmount": fields[16],
                "SupplierName": supplier_name,
                "VendorBranch": vendor_branch,
            }

            records.setdefault(
                invoice_key, []
            ).append(record)

    return records


# ============================================================
# MASTER SUPPLIER FUNCTIONS
# ============================================================

def load_master_supplier(master_path):
    """
    Master Supplier columns:
        Column 3 = VendorTaxId
        Column 4 = SupplierName_Master_EN
        Column 5 = SupplierName_Master_TH
    """

    wb = load_workbook(
        master_path,
        read_only=True,
        data_only=True
    )

    ws = wb.active

    master_data = {}

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        vendor_tax_id = (
            row[2] if len(row) > 2 else None
        )

        supplier_name_en = (
            row[3] if len(row) > 3 else None
        )

        supplier_name_th = (
            row[4] if len(row) > 4 else None
        )

        tax_key = normalize_tax_id(vendor_tax_id)

        if not tax_key:
            continue

        master_data.setdefault(
            tax_key, []
        ).append({
            "SupplierName_Master_EN": clean(
                supplier_name_en
            ),
            "SupplierName_Master_TH": clean(
                supplier_name_th
            ),
        })

    wb.close()

    return master_data


# ============================================================
# E-MESSAGE FUNCTION
# ============================================================

def append_emessage(cell, new_messages):
    """
    Preserve existing Emessage.
    Append new messages without overwriting existing messages.
    Avoid adding identical messages again on rerun.
    """

    existing = clean(cell.value)

    messages = [
        item.strip()
        for item in existing.split(",")
        if item.strip()
    ]

    for message in new_messages:
        if message not in messages:
            messages.append(message)

    cell.value = ",".join(messages)


# ============================================================
# VALIDATION FUNCTION
# ============================================================

def validate_excel(
    excel_path,
    txt_path,
    master_path,
    output_path
):

    # --------------------------------------------------------
    # Check input files
    # --------------------------------------------------------

    for path in (
        excel_path,
        txt_path,
        master_path
    ):
        if not os.path.isfile(path):
            raise FileNotFoundError(
                f"File not found: {path}"
            )

    # --------------------------------------------------------
    # Load TXT H records
    # --------------------------------------------------------

    summary_records = load_summary_h_records(
        txt_path
    )

    # --------------------------------------------------------
    # Load Master Supplier
    # --------------------------------------------------------

    master_records = load_master_supplier(
        master_path
    )

    # --------------------------------------------------------
    # Load OCR Excel
    # --------------------------------------------------------

    wb = load_workbook(excel_path)
    ws = wb.active

    # Map Excel headers.
    column_map = {
        clean(cell.value): cell.column
        for cell in ws[1]
        if cell.value is not None
    }

    required_columns = [
        "TaxInvoiceNo",
        "InvoiceDate",
        "VendorTaxId",
        "VendorBranch",
        "TotalAmount",
        "VATAmount",
        "AmountIncVat",
        "Emessage",
    ]

    missing_columns = [
        col for col in required_columns
        if col not in column_map
    ]

    if missing_columns:
        wb.close()
        raise ValueError(
            "Missing Excel columns: "
            + ", ".join(missing_columns)
        )

    result_count = {
        "PASS": 0,
        "FAIL": 0,
        "TXT_NOT_FOUND": 0,
        "MASTER_NOT_FOUND": 0,
    }

    # ========================================================
    # PROCESS EACH EXCEL ROW
    # ========================================================


    for row_number in range(2, ws.max_row + 1):

        errors = []

        def get_excel_value(column_name):
            return ws.cell(
                row=row_number,
                column=column_map[column_name]
            ).value

        tax_invoice_no = get_excel_value("TaxInvoiceNo")

        if tax_invoice_no is None or clean(tax_invoice_no) == "":
            print(f"Row {row_number}: SKIPPED - Empty TaxInvoiceNo")
            continue
        
        invoice_key = normalize_invoice(
            tax_invoice_no
        )

        txt_record = None

        emessage_cell = ws.cell(
            row=row_number,
            column=column_map["Emessage"]
        )

        # ----------------------------------------------------
        # 1. Find TaxInvoiceNo in Summary TXT
        # ----------------------------------------------------

        tax_invoice_no = get_excel_value(
            "TaxInvoiceNo"
        )

        invoice_key = normalize_invoice(
            tax_invoice_no
        )

        txt_record = None

        if invoice_key:
            matched_records = summary_records.get(
                invoice_key, []
            )

            if matched_records:
                txt_record = matched_records[0]

        # ----------------------------------------------------
        # 2. If TaxInvoiceNo is NOT FOUND:
        #    Skip ALL remaining validation.
        # ----------------------------------------------------

        if txt_record is None:

            append_emessage(
                emessage_cell,
                ["ไม่พบ"]
            )

            result_count["FAIL"] += 1
            result_count["TXT_NOT_FOUND"] += 1

            print(
                f"Row {row_number}: FAIL - ไม่พบ"
            )

            continue

        # ----------------------------------------------------
        # 3. Find VendorTaxId in Master Supplier
        # ----------------------------------------------------

        vendor_tax_id = get_excel_value(
            "VendorTaxId"
        )

        tax_id_key = normalize_tax_id(
            vendor_tax_id
        )

        SupplierName_Master_EN = ""
        SupplierName_Master_TH = ""

        master_found = False

        if tax_id_key:
            master_matches = master_records.get(
                tax_id_key, []
            )

            if master_matches:
                master_found = True

                SupplierName_Master_EN = (
                    master_matches[0][
                        "SupplierName_Master_EN"
                    ]
                )

                SupplierName_Master_TH = (
                    master_matches[0][
                        "SupplierName_Master_TH"
                    ]
                )

        if not master_found:
            errors.append(
                "VendorTaxId ไม่พบใน Master Supplier"
            )

            result_count["MASTER_NOT_FOUND"] += 1

        # ----------------------------------------------------
        # 4. Validate InvoiceDate
        #    Excel vs TXT Field 3
        # ----------------------------------------------------

        excel_date = normalize_date(
            get_excel_value("InvoiceDate")
        )

        txt_date = normalize_date(
            txt_record["InvoiceDate"]
        )

        if excel_date != txt_date:
            errors.append(
                "InvoiceDate ไม่ตรงกัน"
            )

        # ----------------------------------------------------
        # 5. Validate SupplierName
        #    Master Column 4 or 5 vs TXT Field 20
        #    Remove branch in parentheses first.
        # ----------------------------------------------------

        if master_found:

            txt_supplier_name = normalize_text(
                txt_record["SupplierName"]
            )

            master_en = normalize_text(
                SupplierName_Master_EN
            )

            master_th = normalize_text(
                SupplierName_Master_TH
            )

            valid_supplier_names = [
                name for name in (
                    master_en,
                    master_th
                )
                if name
            ]

            if txt_supplier_name not in valid_supplier_names:
                errors.append(
                    "SupplierName ไม่ตรงกัน"
                )

        # ----------------------------------------------------
        # 6. Validate VendorBranch
        #    Excel vs parentheses in TXT Field 20
        # ----------------------------------------------------

        excel_branch = normalize_branch(
            get_excel_value("VendorBranch")
        )

        txt_branch = normalize_branch(
            txt_record["VendorBranch"]
        )

        if excel_branch != txt_branch:
            errors.append(
                "VendorBranch ไม่ตรงกัน"
            )

        # ----------------------------------------------------
        # 7. Validate TotalAmount
        #    Excel vs TXT Field 16
        # ----------------------------------------------------

        excel_total = normalize_amount(
            get_excel_value("TotalAmount")
        )

        txt_total = normalize_amount(
            txt_record["TotalAmount"]
        )

        if excel_total != txt_total:
            errors.append(
                "TotalAmount ไม่ตรงกัน"
            )

        # ----------------------------------------------------
        # 8. Validate VATAmount
        #    Excel vs TXT Field 17
        # ----------------------------------------------------

        excel_vat = normalize_amount(
            get_excel_value("VATAmount")
        )

        txt_vat = normalize_amount(
            txt_record["VATAmount"]
        )

        if excel_vat != txt_vat:
            errors.append(
                "VATAmount ไม่ตรงกัน"
            )

        # ----------------------------------------------------
        # 9. Validate AmountIncVat
        #    Excel vs TXT Field 14
        # ----------------------------------------------------

        excel_amount_inc_vat = normalize_amount(
            get_excel_value("AmountIncVat")
        )

        txt_amount_inc_vat = normalize_amount(
            txt_record["AmountIncVat"]
        )

        if excel_amount_inc_vat != txt_amount_inc_vat:
            errors.append(
                "AmountIncVat ไม่ตรงกัน"
            )

        # ----------------------------------------------------
        # 10. Append all errors to Emessage
        # ----------------------------------------------------

        append_emessage(
            emessage_cell,
            errors
        )

        if errors:
            result_count["FAIL"] += 1
            print(
                f"Row {row_number}: FAIL - "
                + ", ".join(errors)
            )
        else:
            result_count["PASS"] += 1
            print(
                f"Row {row_number}: PASS"
            )

    # ========================================================
    # SAVE OUTPUT
    # ========================================================

    output_dir = os.path.dirname(output_path)

    if output_dir:
        os.makedirs(
            output_dir,
            exist_ok=True
        )

    wb.save(output_path)
    wb.close()

    # ========================================================
    # PRINT SUMMARY
    # ========================================================

    print("\n========== VALIDATION RESULT ==========")
    print(f"PASS             : {result_count['PASS']}")
    print(f"FAIL             : {result_count['FAIL']}")
    print(f"TXT NOT FOUND    : {result_count['TXT_NOT_FOUND']}")
    print(f"MASTER NOT FOUND : {result_count['MASTER_NOT_FOUND']}")
    print(f"\nOutput: {output_path}")


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    validate_excel(
        EXCEL_PATH,
        TEXT_PATH,
        MASTER_PATH,
        OUTPUT_EXCEL_PATH
    )
